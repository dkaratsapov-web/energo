# -*- coding: utf-8 -*-
"""Редактируемый PPTX из той же модели, что и PDF.

Всё — живые объекты PowerPoint: текстовые поля с настоящим текстом,
прямоугольники, линии, кривые, снимки. Строки уже перенесены движком,
перенос в полях выключен, интерлиньяж точный (spcPts), трекинг — spc.
Поэтому PowerPoint показывает ту же вёрстку, что и PDF, а клиент может
править текст.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from .model import Rect, Line, Poly, Dot, Image, Text, Shape, PT
from .render_pdf import _clip_seg
from .layouts import clip_runs
from . import fonts, images

EMU_MM = 36000

# При точном интерлиньяже первая базовая линия стоит на (интерлиньяж − 0,2)
# кегля ниже верха поля — независимо от гарнитуры. Измерено на калибровочной
# колоде через LibreOffice Impress (импорт PPTX эмулирует PowerPoint);
# preflight каждый раз сверяет позиции строк PDF↔PPTX заново.
FIRST_BASELINE = -0.2


def E(mm):
    return Emu(int(round(mm * EMU_MM)))


def _rgb(col):
    return RGBColor(*col)


def _no_line(shape):
    shape.line.fill.background()


def _shadowless(shape):
    """Снимаем стиль темы: иначе PowerPoint/LibreOffice добавят тень и
    обводку по умолчанию. Заливку и линию мы задаём явно."""
    el = shape._element
    st = el.find(qn('p:style'))
    if st is not None:
        el.remove(st)
    sp = el.spPr
    if sp.find(qn('a:effectLst')) is None:
        sp.append(sp.makeelement(qn('a:effectLst'), {}))


def _text(slide, el, ox=0.0, oy=0.0):
    f = fonts.info(el.font)
    n = len(el.lines)
    lead = el.size * el.leading
    wmax = max(el.line_w(i) for i in range(n)) if n else el.w
    w = max(el.w, wmax + 1.0)
    x = el.x if el.align == 'l' else (el.x + el.w - w if el.align == 'r' else el.x + (el.w - w) / 2)
    top = el.baseline(0) - (el.leading + FIRST_BASELINE) * el.size * PT
    h = n * lead * PT + 0.3 * el.size * PT
    tb = slide.shapes.add_textbox(E(x - ox), E(top - oy), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, ln in enumerate(el.lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'c': PP_ALIGN.CENTER}[el.align]
        p.line_spacing = Pt(lead)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = ln
        fn = r.font
        fn.name = f['family']
        fn.size = Pt(el.size)
        fn.bold = f['bold']
        fn.italic = f['italic']
        fn.color.rgb = _rgb(el.color)
        rpr = r._r.get_or_add_rPr()
        rpr.set('spc', str(int(round(el.track * el.size * 100))))
        for tag in ('a:ea', 'a:cs'):
            sub = rpr.find(qn(tag))
            if sub is None:
                sub = rpr.makeelement(qn(tag), {})
                rpr.append(sub)
            sub.set('typeface', f['family'])
    tb.name = f'{el.role}: {el.lines[0][:30] if el.lines else ""}'
    return tb


def _custgeom(s, contours, ox, oy, closed=True, name='shape'):
    """Фигура PowerPoint из контуров — XML пишется напрямую.

    Штатный FreeformBuilder пересчитывает габарит на каждой точке, и знак
    из тысяч точек строится минутами. Здесь — один проход по точкам."""
    from lxml import etree
    xs = [p[0] for c in contours for p in c]
    ys = [p[1] for c in contours for p in c]
    x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
    w, h = max(x1 - x0, 0.01), max(y1 - y0, 0.01)
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x0 - ox), E(y0 - oy), E(w), E(h))
    W_, H_ = int(E(w)), int(E(h))
    parts = [f'<a:custGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
             f'<a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
             f'<a:pathLst><a:path w="{W_}" h="{H_}">']
    for c in contours:
        pts = [(int(E(px - x0)), int(E(py - y0))) for px, py in c]
        parts.append(f'<a:moveTo><a:pt x="{pts[0][0]}" y="{pts[0][1]}"/></a:moveTo>')
        parts.extend(f'<a:lnTo><a:pt x="{x}" y="{y}"/></a:lnTo>' for x, y in pts[1:])
        if closed:
            parts.append('<a:close/>')
    parts.append('</a:path></a:pathLst></a:custGeom>')
    geom = etree.fromstring(''.join(parts))
    sp = sh._element.spPr
    old = sp.find(qn('a:prstGeom'))
    sp.replace(old, geom)
    sh.name = name
    return sh


def render(slides, path, W, H, cache='/tmp/deckkit-img', meta=None, split=False):
    """split=True — каталог: слайд PPTX на каждую полосу, развороты режутся по
    корешку. Иначе — слайд на каждое полотно."""
    fonts.register()
    prs = Presentation()
    blank = prs.slide_layouts[6]
    meta = meta or {}
    cp = prs.core_properties
    cp.title = meta.get('title', '')
    cp.author = meta.get('author', '')
    cp.subject = meta.get('subject', '')
    size = None
    for sl in slides:
        cw, ch = sl.w or W, sl.h or H
        n = len(sl.pages) if split and sl.pages else 1
        vw = cw / n
        if size is None:
            size = (vw, ch)
            prs.slide_width, prs.slide_height = E(vw), E(ch)
        for i in range(n):
            s = prs.slides.add_slide(blank)
            _draw(s, sl, i * vw, 0.0, vw, ch, cache)
            if sl.notes:
                s.notes_slide.notes_text_frame.text = sl.notes
    prs.save(path)
    return path


def _draw(s, sl, ox, oy, vw, vh, cache):
    X0, Y0, X1, Y1 = ox, oy, ox + vw, oy + vh
    for el in sl.items:
        if isinstance(el, Rect):
            x0, y0 = max(el.x, X0), max(el.y, Y0)
            x1, y1 = min(el.x + el.w, X1), min(el.y + el.h, Y1)
            if x1 - x0 < 0.05 or y1 - y0 < 0.05:
                continue
            sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x0 - ox), E(y0 - oy), E(x1 - x0), E(y1 - y0))
            sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(el.fill)
            _no_line(sh); _shadowless(sh)
            sh.name = el.role
        elif isinstance(el, Line):
            seg = _clip_seg(el.x0, el.y0, el.x1, el.y1, X0, Y0, X1, Y1)
            if not seg or abs(seg[0] - seg[2]) + abs(seg[1] - seg[3]) < 0.05:
                continue
            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(seg[0] - ox), E(seg[1] - oy),
                                        E(seg[2] - ox), E(seg[3] - oy))
            ln.line.color.rgb = _rgb(el.color)
            ln.line.width = Pt(el.width)
            _shadowless(ln)
            ln.name = el.role
        elif isinstance(el, Poly):
            runs = [el.pts] if el.closed else clip_runs(el.pts, X0, Y0, X1, Y1)
            for pts in runs:
                sh = _custgeom(s, [pts], ox, oy, el.closed, el.role)
                if el.fill:
                    sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(el.fill)
                else:
                    sh.fill.background()
                if el.stroke:
                    sh.line.color.rgb = _rgb(el.stroke); sh.line.width = Pt(el.width)
                else:
                    _no_line(sh)
                _shadowless(sh)
        elif isinstance(el, Shape):
            bx0, by0, bx1, by1 = el.bbox()
            if bx1 < X0 or bx0 > X1:
                continue
            sh = _custgeom(s, el.contours, ox, oy, True, el.role)
            sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(el.fill)
            _no_line(sh); _shadowless(sh)
        elif isinstance(el, Dot):
            if not (X0 - el.r <= el.cx <= X1 + el.r):
                continue
            sh = s.shapes.add_shape(MSO_SHAPE.OVAL, E(el.cx - el.r - ox), E(el.cy - el.r - oy),
                                    E(2 * el.r), E(2 * el.r))
            sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(el.fill)
            _no_line(sh); _shadowless(sh)
            sh.name = el.role
        elif isinstance(el, Image):
            x0, y0 = max(el.x, X0), max(el.y, Y0)
            x1, y1 = min(el.x + el.w, X1), min(el.y + el.h, Y1)
            if x1 - x0 < 0.05 or y1 - y0 < 0.05:
                continue
            f = images.prepare(el.path, el.w, el.h, el.focus, el.shade, 200, cache)
            pic = s.shapes.add_picture(f, E(x0 - ox), E(y0 - oy), E(x1 - x0), E(y1 - y0))
            # кадр целиком внутри снимка, видимая часть — штатной обрезкой PowerPoint
            pic.crop_left = (x0 - el.x) / el.w
            pic.crop_right = (el.x + el.w - x1) / el.w
            pic.crop_top = (y0 - el.y) / el.h
            pic.crop_bottom = (el.y + el.h - y1) / el.h
            pic.name = 'image'
        elif isinstance(el, Text):
            bx0, by0, bx1, by1 = el.bbox()
            if bx1 < X0 or bx0 > X1:
                continue
            _text(s, el, ox, oy)
